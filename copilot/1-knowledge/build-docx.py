"""K1~K5 지식 문서(.md)를 Copilot Studio 에 올릴 .docx 로 바꾼다.

    python3 copilot/1-knowledge/build-docx.py

필요: python-docx (pip install python-docx). 표·제목·목록·굵게·코드만 다룬다(이 문서들이 쓰는 만큼).
Copilot 지식 검색은 제목과 표를 단서로 쓰므로, 제목 단계와 표 구조를 그대로 살린다.
"""
import pathlib
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "docx"
FONT = "맑은 고딕"


def set_font(doc):
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = doc.styles[name]
        style.font.name = FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.styles["Normal"].font.size = Pt(10)


def add_runs(par, text):
    # **굵게** 와 `코드` 만 살린다
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            par.add_run(part)


def split_row(line):
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip().replace("\\|", "|") for c in cells]


def flush_table(doc, rows):
    header, body = rows[0], [r for r in rows[2:]]
    ncol = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for r, row in enumerate(body, start=1):
        for i in range(ncol):
            cell = table.rows[r].cells[i]
            cell.text = ""
            add_runs(cell.paragraphs[0], row[i] if i < len(row) else "")
    doc.add_paragraph()


def convert(md_path):
    doc = Document()
    set_font(doc)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    table, code = [], None
    for line in lines + [""]:
        if code is not None:
            if line.startswith("```"):
                par = doc.add_paragraph()
                run = par.add_run("\n".join(code))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                code = None
            else:
                code.append(line)
            continue
        if line.startswith("```"):
            code = []
            continue
        if line.lstrip().startswith("|"):
            table.append(split_row(line))
            continue
        if table:
            flush_table(doc, table)
            table = []
        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
            continue
        if not line.strip():
            continue
        if line.startswith(">"):
            par = doc.add_paragraph(style="Intense Quote")
            add_runs(par, line.lstrip("> ").strip())
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)", line)
        number = re.match(r"^(\s*)\d+[.)]\s+(.*)", line)
        if bullet:
            style = "List Bullet 2" if len(bullet.group(1)) >= 2 else "List Bullet"
            add_runs(doc.add_paragraph(style=style), bullet.group(2))
        elif number:
            style = "List Number 2" if len(number.group(1)) >= 2 else "List Number"
            add_runs(doc.add_paragraph(style=style), number.group(2))
        else:
            add_runs(doc.add_paragraph(), line.strip())
    OUT.mkdir(exist_ok=True)
    target = OUT / (md_path.stem + ".docx")
    doc.save(target)
    return target


if __name__ == "__main__":
    for md in sorted(HERE.glob("K*.md")):
        print("만듦:", convert(md).relative_to(HERE.parent.parent))
